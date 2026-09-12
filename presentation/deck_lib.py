"""Minimal slide-description layer that renders the same deck to PPTX
(python-pptx) and PDF (reportlab), so both outputs stay identical.

Units: inches, origin at the top-left of a 13.333 x 7.5 in (16:9) slide.
"""

import re
from dataclasses import dataclass, field
from typing import List, Optional, Union

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph

SLIDE_W = 13.333
SLIDE_H = 7.5

# ---------------------------------------------------------------- palette
NAVY = "0B1F3A"
BLUE = "1F4E8C"
TEAL = "0E9AA7"
AMBER = "B9772E"
TEXT = "1B2430"
MUTED = "5B6B7F"
LINE = "D5DCE5"
PANEL = "F3F6FA"
BLUE_LIGHT = "E8EFF8"
TEAL_LIGHT = "E3F4F6"
AMBER_LIGHT = "FBF1E6"
WHITE = "FFFFFF"

FONT = "Arial"


def blend(c1: str, c2: str, t: float) -> str:
    a = [int(c1[i:i + 2], 16) for i in (0, 2, 4)]
    b = [int(c2[i:i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(a[i] + (b[i] - a[i]) * t):02X}" for i in range(3))


# ---------------------------------------------------------------- model
@dataclass
class Para:
    text: str
    size: Optional[float] = None
    bold: Optional[bool] = None
    color: Optional[str] = None
    bullet: bool = False
    align: Optional[str] = None
    space_after: Optional[float] = None
    italic: bool = False


ParaLike = Union[str, Para]


@dataclass
class Text:
    x: float
    y: float
    w: float
    h: float
    paras: List[ParaLike]
    size: float = 14
    color: str = TEXT
    bold: bool = False
    align: str = "l"
    valign: str = "t"
    inset: float = 0.05
    line_spacing: float = 1.15
    space_after: float = 4


@dataclass
class Rect:
    x: float
    y: float
    w: float
    h: float
    fill: Optional[str] = None
    stroke: Optional[str] = None
    stroke_w: float = 0.75
    radius: float = 0.0
    paras: Optional[List[ParaLike]] = None
    size: float = 12
    color: str = TEXT
    bold: bool = False
    align: str = "c"
    valign: str = "m"
    inset: float = 0.08
    line_spacing: float = 1.1
    space_after: float = 2


@dataclass
class Circle:
    cx: float
    cy: float
    r: float
    fill: Optional[str] = None
    stroke: Optional[str] = None
    stroke_w: float = 0.75
    paras: Optional[List[ParaLike]] = None
    size: float = 12
    color: str = WHITE
    bold: bool = True


@dataclass
class Line:
    x1: float
    y1: float
    x2: float
    y2: float
    color: str = MUTED
    width: float = 1.0
    arrow: bool = False


@dataclass
class Table:
    x: float
    y: float
    w: float
    col_ws: List[float]  # relative weights
    rows: List[List[ParaLike]]
    size: float = 11
    header: bool = True
    header_fill: str = NAVY
    header_color: str = WHITE
    border: str = LINE
    fill: Optional[str] = WHITE
    alt_fill: Optional[str] = None
    first_col_bold: bool = False
    pad: float = 0.07
    min_row_h: float = 0.32
    line_spacing: float = 1.1
    align: Optional[List[str]] = None  # per column: l/c/r
    valign: str = "m"


@dataclass
class Slide:
    bg: str = WHITE
    items: list = field(default_factory=list)

    def add(self, *items):
        self.items.extend(items)
        return self


# ---------------------------------------------------------------- text helpers
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def parse_runs(text: str):
    """'a **b** c' -> [('a ', False), ('b', True), (' c', False)]"""
    runs, pos = [], 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False))
        runs.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], False))
    return runs or [("", False)]


def _as_para(p: ParaLike) -> Para:
    return p if isinstance(p, Para) else Para(p)


# ---------------------------------------------------------------- PDF fonts
_FONTS_READY = False


def _ensure_fonts():
    global _FONTS_READY
    if _FONTS_READY:
        return
    try:
        base = "C:/Windows/Fonts/"
        pdfmetrics.registerFont(TTFont("Arial", base + "arial.ttf"))
        pdfmetrics.registerFont(TTFont("Arial-Bold", base + "arialbd.ttf"))
        pdfmetrics.registerFont(TTFont("Arial-Italic", base + "ariali.ttf"))
        pdfmetrics.registerFont(TTFont("Arial-BoldItalic", base + "arialbi.ttf"))
        pdfmetrics.registerFontFamily(
            "Arial", normal="Arial", bold="Arial-Bold",
            italic="Arial-Italic", boldItalic="Arial-BoldItalic")
        _FONTS_READY = "Arial"
    except Exception:
        _FONTS_READY = "Helvetica"


def _pdf_font():
    _ensure_fonts()
    return _FONTS_READY


def _escape(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _pdf_paragraphs(paras, size, color, bold, align, line_spacing, space_after,
                    italic_default=False):
    """Build reportlab Paragraph objects for a list of paras."""
    font = _pdf_font()
    amap = {"l": TA_LEFT, "c": TA_CENTER, "r": TA_RIGHT}
    out = []
    for raw in paras:
        p = _as_para(raw)
        psize = p.size or size
        pcolor = p.color or color
        pbold = bold if p.bold is None else p.bold
        palign = p.align or align
        psa = space_after if p.space_after is None else p.space_after
        markup = "".join(
            f"<b>{_escape(t)}</b>" if (b or pbold) else _escape(t)
            for t, b in parse_runs(p.text))
        if p.italic:
            markup = f"<i>{markup}</i>"
        style = ParagraphStyle(
            "s", fontName=font, fontSize=psize, leading=psize * 1.2 * line_spacing,
            textColor=HexColor("#" + pcolor), alignment=amap[palign],
            spaceAfter=psa, leftIndent=0.22 * 72 if p.bullet else 0,
            bulletIndent=0.03 * 72 if p.bullet else 0, bulletFontName=font,
            bulletFontSize=psize)
        out.append(Paragraph(markup, style, bulletText="•" if p.bullet else None))
    return out


def measure_paras(paras, width, size, color=TEXT, bold=False, align="l",
                  line_spacing=1.15, space_after=4, inset=0.05) -> float:
    """Height (inches) needed to lay out the paragraphs in a box of `width`."""
    from reportlab.pdfgen.canvas import Canvas
    import io
    c = Canvas(io.BytesIO())
    ps = _pdf_paragraphs(paras, size, color, bold, align, line_spacing, space_after)
    avail = (width - 2 * inset) * 72
    total = 0.0
    for i, p in enumerate(ps):
        _, h = p.wrapOn(c, avail, 10000)
        total += h + (p.style.spaceAfter if i < len(ps) - 1 else 0)
    return total / 72 + 2 * inset


def _draw_paras_pdf(c, x, y, w, h, paras, size, color, bold, align, valign,
                    inset, line_spacing, space_after):
    ps = _pdf_paragraphs(paras, size, color, bold, align, line_spacing, space_after)
    avail_w = (w - 2 * inset) * 72
    heights = []
    for p in ps:
        _, ph = p.wrapOn(c, avail_w, 10000)
        heights.append(ph)
    total = sum(heights) + sum(p.style.spaceAfter for p in ps[:-1])
    box_h = (h - 2 * inset) * 72
    if valign == "m":
        top_off = max(0, (box_h - total) / 2)
    elif valign == "b":
        top_off = max(0, box_h - total)
    else:
        top_off = 0
    cur_top = SLIDE_H * 72 - (y + inset) * 72 - top_off
    for p, ph in zip(ps, heights):
        p.drawOn(c, (x + inset) * 72, cur_top - ph)
        cur_top -= ph + p.style.spaceAfter


# ---------------------------------------------------------------- PPTX helpers
def _rgb(hexs: str) -> RGBColor:
    return RGBColor.from_string(hexs)


def _fill_paras_pptx(tf, paras, size, color, bold, align, valign, inset,
                     line_spacing, space_after):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(inset)
    tf.margin_top = tf.margin_bottom = Inches(inset)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[valign]
    amap = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
    for i, raw in enumerate(paras):
        p = _as_para(raw)
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = amap[p.align or align]
        para.line_spacing = line_spacing
        para.space_after = Pt(space_after if p.space_after is None else p.space_after)
        psize = p.size or size
        pcolor = p.color or color
        pbold = bold if p.bold is None else p.bold
        for t, b in parse_runs(p.text):
            r = para.add_run()
            r.text = t
            r.font.size = Pt(psize)
            r.font.bold = bool(b or pbold)
            r.font.italic = p.italic
            r.font.name = FONT
            r.font.color.rgb = _rgb(pcolor)
        if p.bullet:
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.22)))
            pPr.set("indent", str(-Inches(0.19)))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")


def _shape_basic(shape, fill, stroke, stroke_w):
    shape.shadow.inherit = False
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
    else:
        shape.fill.background()
    if stroke:
        shape.line.color.rgb = _rgb(stroke)
        shape.line.width = Pt(stroke_w)
    else:
        shape.line.fill.background()


def _add_rect_pptx(slide, r: Rect):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if r.radius > 0 else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(r.x), Inches(r.y), Inches(r.w), Inches(r.h))
    if r.radius > 0:
        s.adjustments[0] = min(0.5, r.radius / min(r.w, r.h))
    _shape_basic(s, r.fill, r.stroke, r.stroke_w)
    if r.paras:
        _fill_paras_pptx(s.text_frame, r.paras, r.size, r.color, r.bold, r.align,
                         r.valign, r.inset, r.line_spacing, r.space_after)
    else:
        s.text_frame.text = ""
    return s


def _add_circle_pptx(slide, c: Circle):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(c.cx - c.r), Inches(c.cy - c.r),
                               Inches(2 * c.r), Inches(2 * c.r))
    _shape_basic(s, c.fill, c.stroke, c.stroke_w)
    if c.paras:
        _fill_paras_pptx(s.text_frame, c.paras, c.size, c.color, c.bold, "c", "m",
                         0.02, 1.0, 0)


def _add_text_pptx(slide, t: Text):
    tb = slide.shapes.add_textbox(Inches(t.x), Inches(t.y), Inches(t.w), Inches(t.h))
    _fill_paras_pptx(tb.text_frame, t.paras, t.size, t.color, t.bold, t.align,
                     t.valign, t.inset, t.line_spacing, t.space_after)


def _add_line_pptx(slide, ln: Line):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ln.x1), Inches(ln.y1),
                                      Inches(ln.x2), Inches(ln.y2))
    conn.line.color.rgb = _rgb(ln.color)
    conn.line.width = Pt(ln.width)
    if ln.arrow:
        lnx = conn.line._get_or_add_ln()
        tail = etree.SubElement(lnx, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")


def _set_cell_border(cell, color, width_pt):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag))
        ln.set("w", str(int(width_pt * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(sf, qn("a:srgbClr"))
        clr.set("val", color)
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
        tcPr.insert(i, ln)


def _table_layout(t: Table):
    total = sum(t.col_ws)
    widths = [t.w * cw / total for cw in t.col_ws]
    heights = []
    for ri, row in enumerate(t.rows):
        is_header = t.header and ri == 0
        h = t.min_row_h
        for ci, cell in enumerate(row):
            bold = is_header or (t.first_col_bold and ci == 0)
            al = (t.align[ci] if t.align else "l")
            h = max(h, measure_paras([cell], widths[ci], t.size, bold=bold, align=al,
                                     line_spacing=t.line_spacing, space_after=2,
                                     inset=t.pad))
        heights.append(h)
    return widths, heights


def _add_table_pptx(slide, t: Table):
    widths, heights = _table_layout(t)
    nrows, ncols = len(t.rows), len(t.col_ws)
    gf = slide.shapes.add_table(nrows, ncols, Inches(t.x), Inches(t.y), Inches(t.w),
                                Inches(sum(heights)))
    tbl = gf.table
    # plain style: no banding, no theme colouring
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    style_id = tblPr.find(qn("a:tableStyleId"))
    if style_id is None:
        style_id = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
    for ci, w in enumerate(widths):
        tbl.columns[ci].width = Inches(w)
    for ri, h in enumerate(heights):
        tbl.rows[ri].height = Inches(h)
    for ri, row in enumerate(t.rows):
        is_header = t.header and ri == 0
        for ci, raw in enumerate(row):
            cell = tbl.cell(ri, ci)
            fill = t.header_fill if is_header else (
                t.alt_fill if (t.alt_fill and ri % 2 == 0) else t.fill)
            if fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(fill)
            else:
                cell.fill.background()
            _set_cell_border(cell, t.border, 0.75)
            cell.margin_left = cell.margin_right = Inches(t.pad)
            cell.margin_top = cell.margin_bottom = Inches(t.pad)
            cell.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                                    "b": MSO_ANCHOR.BOTTOM}[t.valign]
            color = t.header_color if is_header else TEXT
            bold = is_header or (t.first_col_bold and ci == 0)
            al = (t.align[ci] if t.align else "l")
            tf = cell.text_frame
            tf.word_wrap = True
            amap = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
            p = _as_para(raw)
            para = tf.paragraphs[0]
            para.alignment = amap[p.align or al]
            para.line_spacing = t.line_spacing
            for txt, b in parse_runs(p.text):
                r = para.add_run()
                r.text = txt
                r.font.size = Pt(p.size or t.size)
                r.font.bold = bool(b or (bold if p.bold is None else p.bold))
                r.font.name = FONT
                r.font.color.rgb = _rgb(p.color or color)


# ---------------------------------------------------------------- PDF drawing
def _pdf_rect(c, r: Rect):
    H = SLIDE_H * 72
    x, y, w, h = r.x * 72, H - (r.y + r.h) * 72, r.w * 72, r.h * 72
    if r.fill:
        c.setFillColor(HexColor("#" + r.fill))
    if r.stroke:
        c.setStrokeColor(HexColor("#" + r.stroke))
        c.setLineWidth(r.stroke_w)
    if r.fill or r.stroke:
        if r.radius > 0:
            c.roundRect(x, y, w, h, r.radius * 72, stroke=bool(r.stroke), fill=bool(r.fill))
        else:
            c.rect(x, y, w, h, stroke=bool(r.stroke), fill=bool(r.fill))
    if r.paras:
        _draw_paras_pdf(c, r.x, r.y, r.w, r.h, r.paras, r.size, r.color, r.bold,
                        r.align, r.valign, r.inset, r.line_spacing, r.space_after)


def _pdf_circle(c, ci: Circle):
    H = SLIDE_H * 72
    if ci.fill:
        c.setFillColor(HexColor("#" + ci.fill))
    if ci.stroke:
        c.setStrokeColor(HexColor("#" + ci.stroke))
        c.setLineWidth(ci.stroke_w)
    c.circle(ci.cx * 72, H - ci.cy * 72, ci.r * 72, stroke=bool(ci.stroke), fill=bool(ci.fill))
    if ci.paras:
        _draw_paras_pdf(c, ci.cx - ci.r, ci.cy - ci.r, 2 * ci.r, 2 * ci.r, ci.paras,
                        ci.size, ci.color, ci.bold, "c", "m", 0.02, 1.0, 0)


def _pdf_line(c, ln: Line):
    import math
    H = SLIDE_H * 72
    x1, y1, x2, y2 = ln.x1 * 72, H - ln.y1 * 72, ln.x2 * 72, H - ln.y2 * 72
    c.setStrokeColor(HexColor("#" + ln.color))
    c.setFillColor(HexColor("#" + ln.color))
    c.setLineWidth(ln.width)
    ang = math.atan2(y2 - y1, x2 - x1)
    if ln.arrow:
        size = 5 + ln.width * 2.2
        bx, by = x2 - size * math.cos(ang), y2 - size * math.sin(ang)
        c.line(x1, y1, bx, by)
        p = c.beginPath()
        p.moveTo(x2, y2)
        p.lineTo(bx + size * 0.5 * math.sin(ang), by - size * 0.5 * math.cos(ang))
        p.lineTo(bx - size * 0.5 * math.sin(ang), by + size * 0.5 * math.cos(ang))
        p.close()
        c.drawPath(p, stroke=0, fill=1)
    else:
        c.line(x1, y1, x2, y2)


def _pdf_table(c, t: Table):
    widths, heights = _table_layout(t)
    y = t.y
    for ri, row in enumerate(t.rows):
        is_header = t.header and ri == 0
        x = t.x
        for ci, raw in enumerate(row):
            fill = t.header_fill if is_header else (
                t.alt_fill if (t.alt_fill and ri % 2 == 0) else t.fill)
            color = t.header_color if is_header else TEXT
            bold = is_header or (t.first_col_bold and ci == 0)
            al = (t.align[ci] if t.align else "l")
            _pdf_rect(c, Rect(x, y, widths[ci], heights[ri], fill=fill, stroke=t.border,
                              stroke_w=0.75, paras=[raw], size=t.size, color=color,
                              bold=bold, align=al, valign=t.valign, inset=t.pad,
                              line_spacing=t.line_spacing, space_after=2))
            x += widths[ci]
        y += heights[ri]


# ---------------------------------------------------------------- renderers
def render_pptx(slides: List[Slide], path: str):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    for sd in slides:
        s = prs.slides.add_slide(blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(sd.bg)
        for it in sd.items:
            if isinstance(it, Rect):
                _add_rect_pptx(s, it)
            elif isinstance(it, Circle):
                _add_circle_pptx(s, it)
            elif isinstance(it, Text):
                _add_text_pptx(s, it)
            elif isinstance(it, Line):
                _add_line_pptx(s, it)
            elif isinstance(it, Table):
                _add_table_pptx(s, it)
    prs.save(path)


def render_pdf(slides: List[Slide], path: str, title: str = ""):
    _ensure_fonts()
    c = rl_canvas.Canvas(path, pagesize=(SLIDE_W * 72, SLIDE_H * 72))
    c.setTitle(title)
    for sd in slides:
        c.setFillColor(HexColor("#" + sd.bg))
        c.rect(0, 0, SLIDE_W * 72, SLIDE_H * 72, stroke=0, fill=1)
        for it in sd.items:
            if isinstance(it, Rect):
                _pdf_rect(c, it)
            elif isinstance(it, Circle):
                _pdf_circle(c, it)
            elif isinstance(it, Text):
                _draw_paras_pdf(c, it.x, it.y, it.w, it.h, it.paras, it.size, it.color,
                                it.bold, it.align, it.valign, it.inset, it.line_spacing,
                                it.space_after)
            elif isinstance(it, Line):
                _pdf_line(c, it)
            elif isinstance(it, Table):
                _pdf_table(c, it)
        c.showPage()
    c.save()
